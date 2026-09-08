from datetime import date, datetime
from pathlib import Path
import argparse
import re
import sqlite3
import unicodedata

from openpyxl import load_workbook

from app import DB, FIELDS, init_db


def normalize(value):
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", text).strip().lower()


def text_value(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def column_mapping(headers):
    labels = {
        normalize("Nome do Assistido(a)"): "nome",
        normalize("Nome Social"): "nome_social",
        normalize("CPF"): "cpf",
        normalize("RG/órgão emissor"): "rg",
        normalize("Data de nascimento"): "data_nascimento",
        normalize("Nome da mãe"): "nome_mae",
        normalize("Número do Processo"): "processo",
        normalize("Órgão Judicial/Vara"): "vara",
        normalize("Nº Inscrição (ID único)"): "rji",
        normalize("Telefone(s) WhatsApp"): "telefone",
        normalize("Data de Atendimento"): "data_atendimento",
        normalize("Raça/cor da pele (Declarada)"): "raca",
        normalize("Sexo (biológico)"): "sexo",
        normalize("Identidade de Gênero (Declarada)"): "identidade_genero",
        normalize("Orientação Sexual (autodeclaração)"): "orientacao_sexual",
        normalize("Escolaridade"): "escolaridade",
        normalize("Pessoa com Deficiência"): "pcd",
        normalize("Tipo de Deficiência"): "tipo_deficiencia",
        normalize("Nacionalidade"): "nacionalidade",
        normalize("País"): "pais",
        normalize("Ocupação"): "ocupacao",
        normalize("Profissão"): "profissao",
        normalize("Religião"): "religiao",
        normalize("Diploma Legal"): "diploma_legal",
        normalize("Artigo/Capitulação"): "artigo",
        normalize("Tipo Penal"): "tipo_penal",
        normalize("Grupamento Penal"): "grupamento_penal",
        normalize("Natureza do Atendimento"): "natureza_atendimento",
        normalize("Tipo de Medida/Alternativa"): "medida",
        normalize("Grupo de Responsabilização"): "grupo_responsabilizacao",
        normalize("Status do Atendimento"): "status",
        normalize("Atendimento Individual (Assistente Social, Psicólogo, Jurídico)"): "atendimento_individual",
        normalize("Comparecimento Voluntário"): "comparecimento",
        normalize("Data de Término da Medida"): "termino_medida",
        normalize("Situação Final"): "situacao_final",
        normalize("Observação/Justificativas"): "observacoes",
    }
    return {index: labels.get(normalize(header)) for index, header in enumerate(headers)}


def is_duplicate(connection, values):
    for field in ("cpf", "processo", "rji"):
        value = values.get(field, "")
        if value and connection.execute(
            f"SELECT 1 FROM pessoas WHERE {field} = ? LIMIT 1", (value,)
        ).fetchone():
            return True
    return False


def import_rows(path, sheet_name):
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Aba não encontrada: {sheet_name}")
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = next(
        (
            row
            for row in rows
            if any(normalize(value) == normalize("Nome do Assistido(a)") for value in row)
        ),
        None,
    )
    if headers is None:
        raise ValueError("Cabeçalho da base não encontrado")
    mapping = column_mapping(headers)
    missing = [field for field in FIELDS if field not in mapping.values()]
    if missing:
        raise ValueError(f"Colunas não mapeadas: {', '.join(missing)}")

    init_db()
    connection = sqlite3.connect(DB)
    inserted = skipped_empty = skipped_duplicate = 0
    now = datetime.now().isoformat(timespec="seconds")
    placeholders = ", ".join("?" for _ in FIELDS)
    columns = ", ".join(FIELDS)
    try:
        for row in rows:
            values = {field: text_value(row[index]) for index, field in mapping.items() if field}
            if not values.get("nome"):
                skipped_empty += 1
                continue
            if is_duplicate(connection, values):
                skipped_duplicate += 1
                continue
            connection.execute(
                f"INSERT INTO pessoas(criado_em, criado_por, {columns}) VALUES (?, ?, {placeholders})",
                (now, None, *(values.get(field, "") for field in FIELDS)),
            )
            inserted += 1
        connection.commit()
    finally:
        connection.close()
        workbook.close()
    return inserted, skipped_empty, skipped_duplicate


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Importa a base de atendimentos para o CIAP.")
    parser.add_argument("planilha", type=Path)
    parser.add_argument("--aba", default="Dash_Base_Dados")
    args = parser.parse_args()
    result = import_rows(args.planilha, args.aba)
    print(f"inseridos={result[0]} vazios={result[1]} duplicados={result[2]}")